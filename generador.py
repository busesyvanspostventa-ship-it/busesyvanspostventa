from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps
from anexos import validar_anexos, paginas_pdf

CAMPOS = ('modelo', 'fecha', 'vin', 'patente', 'sistema', 'motivo',
          'entidades', 'analisis', 'numero_informe',
          'cliente', 'km', 'sucursal', 'tipo_atencion')


def generar_informe(datos, fotos=(), plantilla=None, foto_vin=None, anexos=()):
    """Genera el Word en memoria. Las fotografías no se almacenan."""
    if not foto_vin:
        raise ValueError('Es obligatorio adjuntar la fotografía del VIN.')
    validar_anexos(anexos)
    fotos = [(foto_vin, 'Evidencia obligatoria: fotografía del VIN del vehículo. VIN informado: '
              + str(datos.get('vin') or 'No informado'))] + list(fotos)
    doc = Document(plantilla or Path(__file__).with_name('plantilla_foton.docx'))
    valores = {k: str(datos.get(k, '') or '').strip() or 'No informado' for k in CAMPOS}

    # Retira el bloque anterior para colocar una única firma al final.
    for p in list(doc.paragraphs):
        if p.text.strip().upper() == 'RESPONSABLE DEL INFORME' or p.text.strip() in (
            '{{responsable}}', '{{cargo}}', '{{correo}}'
        ):
            p._p.getparent().remove(p._p)

    # Reemplaza marcadores de texto del documento y de sus tablas.
    parrafos = list(doc.paragraphs)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                parrafos.extend(celda.paragraphs)

    tiene_numero = any('{{numero_informe}}' in p.text for p in parrafos)
    for p in parrafos:
        texto = p.text
        for k, v in valores.items():
            texto = texto.replace('{{' + k + '}}', v)
        # Evita duplicar responsable en otra posición: siempre va al final.
        for k in ('responsable', 'cargo', 'correo'):
            texto = texto.replace('{{' + k + '}}', '')
        if texto != p.text:
            if p.runs:
                p.runs[0].text = texto
                for r in p.runs[1:]:
                    r.text = ''
            else:
                p.add_run(texto)

    ancla = next((p for p in doc.paragraphs if '{{fotos}}' in p.text), None)
    if ancla is None:
        # Permite trabajar con plantillas sin el marcador de fotografías.
        ancla = doc.add_paragraph('{{fotos}}')

    def insertar_antes_ancla():
        nodo = OxmlElement('w:p')
        ancla._p.addprevious(nodo)
        p = Paragraph(nodo, ancla._parent)
        p.paragraph_format.space_after = Pt(6)
        return p

    if not fotos:
        insertar_antes_ancla().add_run('Sin evidencia fotográfica adjunta.')

    for i, (contenido, descripcion) in enumerate(fotos, 1):
        # Todas las fotos ocupan el mismo recuadro de 14 x 8 cm.
        # Se agrega margen blanco si hace falta, sin recortar ni deformar.
        with Image.open(BytesIO(contenido)) as original:
            orientada = ImageOps.exif_transpose(original).convert('RGBA')
            fondo = Image.new('RGBA', orientada.size, 'white')
            fondo.alpha_composite(orientada)
            im = fondo.convert('RGB')
            im.thumbnail((1400, 800), Image.Resampling.LANCZOS)
            marco = Image.new('RGB', (1400, 800), 'white')
            marco.paste(im, ((1400 - im.width) // 2, (800 - im.height) // 2))
            buffer = BytesIO()
            marco.save(buffer, format='JPEG', quality=92)
        buffer.seek(0)
        p = insertar_antes_ancla()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        p.paragraph_format.line_spacing = 1
        p.add_run().add_picture(buffer, width=Cm(14), height=Cm(8))
        p = insertar_antes_ancla()
        p.paragraph_format.keep_together = True
        p.add_run(f'Fotografía {i}. ').bold = True
        p.add_run(str(descripcion).strip() or 'Sin descripción.')

    ancla._p.getparent().remove(ancla._p)

    numero = str(datos.get('numero_informe', '') or '').strip()
    if numero and not tiene_numero:
        nodo = OxmlElement('w:p')
        doc.paragraphs[0]._p.addnext(nodo)
        p = Paragraph(nodo, doc._body)
        p.alignment = 2
        p.add_run(f'Informe N.º {numero}').bold = True

    # Responsable centrado al final, sin encabezado.
    parrafos_firma = []
    if str(datos.get('responsable', '') or '').strip():
        for clave in ('responsable', 'cargo', 'correo'):
            p = doc.add_paragraph(str(datos.get(clave, '') or '').strip())
            parrafos_firma.append(p._p)
            if clave == 'responsable':
                p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = clave != 'correo'
            for r in p.runs:
                r.font.size = Pt(12)
                r.bold = clave == 'responsable'

    # Los anexos van después de la firma. Cada página PDF ocupa una página Word.
    for indice, (nombre, contenido) in enumerate(anexos, 1):
        for pagina, total, imagen, ancho, alto in paginas_pdf(contenido):
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(6)
            p.add_run(f'Anexo {indice} · {nombre[:100]} · página {pagina}/{total}').bold = True
            seccion = doc.sections[-1]
            ancho_max = seccion.page_width - seccion.left_margin - seccion.right_margin
            alto_max = seccion.page_height - seccion.top_margin - seccion.bottom_margin - Cm(2.5)
            escala = min(ancho_max / ancho, alto_max / alto)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            p.add_run().add_picture(BytesIO(imagen), width=int(ancho * escala), height=int(alto * escala))

    # Tipografía uniforme; los párrafos de contenido quedan justificados.
    def formatear(parrafo):
        contiene_imagen = bool(parrafo._p.xpath('.//w:drawing'))
        es_titulo = parrafo.text.strip().casefold() == 'informe técnico de falla'
        es_firma = parrafo._p in parrafos_firma
        if es_titulo or es_firma:
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.paragraph_format.left_indent = Cm(0)
            parrafo.paragraph_format.right_indent = Cm(0)
            parrafo.paragraph_format.first_line_indent = Cm(0)
        elif not contiene_imagen and not parrafo.text.startswith('Informe N.º'):
            parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in parrafo.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(14 if es_titulo else 12)
            if es_titulo:
                run.bold = True
            fuentes = run._element.get_or_add_rPr().rFonts
            if fuentes is not None:
                for atributo in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    fuentes.set(qn('w:' + atributo), 'Calibri')

    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(12)
    for p in doc.paragraphs:
        formatear(p)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    formatear(p)

    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue()
